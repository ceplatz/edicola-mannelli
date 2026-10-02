#!/usr/bin/env python3
"""Converte il foglio di un distributore (layout PIER: schede PIER / CARICA / SCARICA)
in file CSV puliti per l'importatore dell'app (Impostazioni -> Importa).

Uso:  python3 tools/convert_workbook.py FILE.xlsx "Nome distributore" CARTELLA_USCITA

NON salvare i file generati nel repository: contengono i prezzi di acquisto.
"""
import csv, re, sys, os, datetime
import openpyxl

SMALL = {'di','del','della','dei','delle','e','a','da','in','con','per','le','la','il','lo','un','una','al','alla'}
CAT_FIX = {'MNIATURA': 'MINIATURE', 'MINIATURA': 'MINIATURE', 'MONETA': 'MONETE'}

def clean_name(s):
    s = re.sub(r'\s+', ' ', str(s)).strip()
    if len(s) > 3 and s.isupper():
        words = s.title().split(' ')
        s = ' '.join(w.lower() if (i and w.lower() in SMALL) else w for i, w in enumerate(words))
    elif s and s[0].islower():
        s = s[0].upper() + s[1:]
    return s

def clean_cat(s):
    s = re.sub(r'\s+', ' ', str(s or '')).strip().upper()
    words = CAT_FIX.get(s, s).title().split(' ')
    return ' '.join(w.lower() if (i and w.lower() in SMALL) else w for i, w in enumerate(words))

def clean_supplier_code(s):
    if s is None: return ''
    s = re.sub(r'\s+', ' ', str(s)).strip()
    m = re.fullmatch(r'(?i)art\.?\s*(\d+)', s)
    return f'Art.{m.group(1)}' if m else s

def num(v, default=0.0):
    try: return float(v)
    except (TypeError, ValueError): return default

def main(xlsx, distributore, out):
    os.makedirs(out, exist_ok=True)
    wb = openpyxl.load_workbook(xlsx, data_only=True)      # valori calcolati dei fogli
    ws = wb['PIER'] if 'PIER' in wb.sheetnames else wb.worksheets[0]
    prodotti, flags, skipped = [], [], []
    code_by_lower = {}
    for r in ws.iter_rows(min_row=4, values_only=True):
        code, forn, cat, desc = r[0], r[1], r[2], r[3]
        if code is None and desc is None: continue
        if isinstance(code, str) and (code.startswith('▶') or code.upper().startswith('VALORE TOTALE')): continue
        code = str(code).strip() if code is not None else ''
        if code.upper() == 'M85':
            skipped.append((code, str(desc), 'segnaposto "generico" (non è un articolo)')); continue
        if desc is None or not str(desc).strip():
            skipped.append((code, '', 'codice senza descrizione (articolo non compilato), saltato')); continue
        nome = clean_name(desc)
        stock = num(r[9]); pa = num(r[5]); pv = num(r[6]); smin = int(num(r[11]))
        qta = max(int(round(stock)), 0)
        ricontare = stock < 0
        row = dict(distributore=distributore, codice=code.upper(), codice_fornitore=clean_supplier_code(forn),
                   categoria=clean_cat(cat), nome=nome, prezzo_acquisto=f'{pa:.2f}', prezzo_vendita=f'{pv:.2f}',
                   scorta_min=smin, qta_magazzino=qta, qta_negozio=0,
                   da_ricontare='si' if ricontare else '',
                   nota_ricontare=f'Il foglio indicava {int(round(stock))}' if ricontare else '')
        prodotti.append(row)
        if code: code_by_lower[code.lower()] = code.upper()
        if ricontare: flags.append((code, nome, 'Giacenza negativa nel foglio', int(round(stock)), 'Ricontare i pezzi e correggere la quantità'))
        if pv == 0: flags.append((code, nome, 'Prezzo di vendita mancante (0)', '0', 'Inserire il prezzo di vendita'))
        if not code: flags.append(('', nome, 'Articolo senza codice', '', 'Assegnare un codice (es. P1, M1)'))
    # nomi uguali dopo la pulizia (es. 3 articoli "Moneta")
    seen = {}
    for p in prodotti: seen.setdefault(p['nome'].lower(), []).append(p)
    for n, ps in seen.items():
        if len(ps) > 1:
            for p in ps: flags.append((p['codice'], p['nome'], 'Stesso nome di altri articoli', ', '.join(x['codice'] or '?' for x in ps), 'Distinguere il nome (es. misura/colore)'))

    def history(sheet, tipo, prefix):
        w = wb[sheet]; rows = []; cur = None; n = 0
        for r in w.iter_rows(min_row=3, values_only=True):
            if isinstance(r[1], (datetime.datetime, datetime.date)): cur = r[1].date() if hasattr(r[1], 'date') else r[1]
            code = (str(r[2]).strip() if r[2] is not None else '')
            if not code or r[4] is None: continue
            n += 1
            rows.append(dict(id_riga=f'{prefix}-{n:04d}', distributore=distributore, data=cur.isoformat() if cur else '',
                             tipo=tipo, codice=code, qta=int(num(r[4])), ndoc=(r[0] or ''), nota=(r[5] or '')))
        return rows
    carichi = history('CARICA', 'carico', 'c'); scarichi = history('SCARICA', 'vendita', 's')
    hist = []
    r7 = [h for h in carichi + scarichi if h['codice'].lower() == 'r7']
    for h in carichi + scarichi:
        low = h['codice'].lower()
        if low == 'r7':
            skipped.append((h['codice'], h['tipo'], f"{h['qta']} pz il {h['data']}: carico e vendita di 134 pz sullo stesso articolo (probabile rettifica), saltati"))
            continue
        if low not in code_by_lower:
            skipped.append((h['codice'], h['tipo'], f"{h['qta']} pz il {h['data']}: codice non presente nell'elenco articoli")); continue
        h['codice'] = code_by_lower[low]; hist.append(h)

    def write(name, fields, rows):
        with open(os.path.join(out, name), 'w', newline='', encoding='utf-8') as f:
            w = csv.DictWriter(f, fieldnames=fields); w.writeheader(); w.writerows(rows)
    write('prodotti.csv', ['distributore','codice','codice_fornitore','categoria','nome','prezzo_acquisto','prezzo_vendita','scorta_min','qta_magazzino','qta_negozio','da_ricontare','nota_ricontare'], prodotti)
    write('storico.csv', ['id_riga','distributore','data','tipo','codice','qta','ndoc','nota'], hist)

    # elenco "da controllare" per Federica
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    nb = Workbook(); s = nb.active; s.title = 'Da controllare'
    s.append(['Codice', 'Articolo', 'Problema', 'Valore nel foglio', 'Cosa fare'])
    for fl in sorted(flags, key=lambda x: (x[2], x[0])): s.append(list(fl))
    s2 = nb.create_sheet('Righe saltate'); s2.append(['Codice', 'Riga', 'Motivo'])
    for sk in skipped: s2.append(list(sk))
    for sh in (s, s2):
        for c in sh[1]: c.font = Font(name='Arial', bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='3A2E28')
        for row in sh.iter_rows(min_row=2):
            for c in row: c.font = Font(name='Arial'); c.alignment = Alignment(wrap_text=True, vertical='top')
        for i, wd in enumerate([12, 36, 38, 22, 44][:sh.max_column], 1): sh.column_dimensions[get_column_letter(i)].width = wd
        sh.freeze_panes = 'A2'
    nb.save(os.path.join(out, 'DA-CONTROLLARE.xlsx'))
    print(f'prodotti.csv: {len(prodotti)} | storico.csv: {len(hist)} (carichi {sum(h["tipo"]=="carico" for h in hist)}, vendite {sum(h["tipo"]=="vendita" for h in hist)}) | da controllare: {len(flags)} | saltate: {len(skipped)}')

if __name__ == '__main__':
    if len(sys.argv) != 4: sys.exit(__doc__)
    main(*sys.argv[1:])
